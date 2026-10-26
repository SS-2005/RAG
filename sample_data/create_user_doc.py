import docx
import os

def generate_user_doc():
    doc = docx.Document()
    
    # Title & Front Matter
    doc.add_heading("A PROJECT SYNOPSIS On CYBER ARENA", level=0)
    
    doc.add_paragraph("Submitted By:\nPranav Tembhare (70)\nShivam Tiwari (72)\nSahil Shaikh (64)\nRahul Bhagwatkar (04)")
    doc.add_paragraph("Under the supervision of:\nProf. Anuja Chandane")
    doc.add_paragraph("Department of Computer Science & Engineering in Artificial Intelligence and Machine Learning\nSaraswati College of Engineering, Kharghar, Navi Mumbai (Affiliated to University of Mumbai) 2026-27")
    
    # Certificate
    doc.add_heading("CERTIFICATE", level=1)
    doc.add_paragraph(
        "This is to certify that the requirements for the synopsis entitled 'CYBER ARENA' have been successfully completed by the following students:\n"
        "Roll numbers: 70 - Pranav Tembhare, 72 - Shivam Tiwari, 64 - Sahil Shaikh, 04 - Rahul Bhagwatkar.\n"
        "In partial fulfillment of Sem-VII, Bachelor of Engineering of Mumbai University.\n"
        "Internal Guide: Prof. Anuja Chandane, Project Coordinator: Prof. Snehal Gavale, Head of Department: Prof. Anuja Chandane, Principal: Dr. Manjusha Deshmukh."
    )
    
    # Abstract
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "Cybersecurity testing has traditionally relied on predefined scripts and manually executed commands. "
        "CyberArena is a modular, browser-based cybersecurity simulation platform designed to study and evaluate LLM-powered autonomous agents in controlled virtual environments. "
        "Through a React-based interface, a user configures an experiment by selecting virtual machines (VMs), assigning a locally hosted LLM served through Ollama, and defining each agent's role and objective. "
        "A FastAPI backend and an orchestrator then start one agent per VM. Each agent follows an iterative workflow: it observes its environment, reasons about the next step, selects a tool from a controlled registry, executes it inside the VM, and analyzes the feedback."
    )
    
    # Scope
    doc.add_heading("1.4 Scope of the Study", level=1)
    doc.add_paragraph(
        "In scope: The project will develop a modular cyber-range platform with the following components:\n"
        "- isolated VirtualBox virtual machines as the target environments;\n"
        "- an Ollama-based provider for locally hosted open-source LLMs;\n"
        "- a custom agent runtime implementing the observe-reason-act loop;\n"
        "- a controlled tool registry and execution layer;\n"
        "- an evaluation and logging engine backed by SQLite;\n"
        "- a React-based browser dashboard on a FastAPI backend.\n"
        "The initial implementation will use lightweight local models and will handle two simultaneous virtual machines, each controlled by its own agent."
    )
    
    # System Architecture & Requirements
    doc.add_heading("4.2 Design Details", level=1)
    doc.add_paragraph(
        "The architecture connects a React browser interface to a FastAPI backend. The orchestrator spawns agents that use an Ollama LLM provider and interface with VirtualBox environments via a controlled tool registry. "
        "The orchestrator runs two agents concurrently, each in its own isolated VM."
    )
    
    doc.add_heading("Chapter 6 Software and Hardware Requirements", level=1)
    doc.add_paragraph(
        "Software Requirements:\n"
        "- Backend: Python, FastAPI\n"
        "- Frontend: React, TypeScript\n"
        "- LLM runtime: Ollama\n"
        "- VM platform: Oracle VirtualBox\n"
        "- Database: SQLite\n"
        "- Tests: Pytest"
    )
    doc.add_paragraph(
        "Hardware Requirements:\n"
        "- Memory: Minimum 8 GB RAM (16 GB recommended when running two VMs and an LLM together)\n"
        "- Storage: SSD\n"
        "- Processor: Multi-core CPU with hardware virtualization enabled"
    )
    
    # Acknowledgement
    doc.add_heading("Acknowledgement", level=1)
    doc.add_paragraph(
        "We express our gratitude to Prof. Anuja Chandane as project guide, Prof. Snehal Gavale as coordinator, Dr. Manjusha Deshmukh as principal.\n"
        "Students:\nPranav Tembhare (70)\nShivam Tiwari (72)\nSahil Shaikh (64)\nRahul Bhagwatkar (04)"
    )

    out_path = "sample_data/cyber_arena_synopsis.docx"
    doc.save(out_path)
    print("Saved to", out_path)

if __name__ == "__main__":
    generate_user_doc()
